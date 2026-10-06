import math, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Four_Corners_Commercial_Quote_Template.xlsx")

F = "Arial"
NAVY = "1F3864"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="008000")
f_hdr = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title = Font(name=F, size=16, bold=True, color=NAVY)
f_sec = Font(name=F, size=11, bold=True, color=NAVY)
f_ex = Font(name=F, size=10, color="0000FF", italic=True)
fill_hdr = PatternFill("solid", fgColor=NAVY)
fill_in = PatternFill("solid", fgColor="FFF9C4")
fill_band = PatternFill("solid", fgColor="DCE3EF")
fill_tot = PatternFill("solid", fgColor="E7E6E6")
thin = Side(style="thin", color="A6A6A6")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
topline = Border(top=Side(style="medium", color=NAVY))
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
CUR = '$#,##0.00;($#,##0.00);"-"'
CUR0 = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
QTY = '#,##0.00;(#,##0.00);"-"'

wb = Workbook()

def hdr(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = f_hdr, fill_hdr, center, box

def inp(c, v=None, fmt=None):
    if v is not None:
        c.value = v
    c.font, c.fill, c.border = f_input, fill_in, box
    if fmt:
        c.number_format = fmt

def setfont(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)

# ---------------------------------------------------------------- Lists
GROUPS = ["Base", "Alt 1", "Alt 2", "Alt 3", "Alt 4", "Alt 5", "Optional", "Excluded"]
SECTIONS = [
    "01 52 00 Temporary Storage",
    "01 54 16 Temporary Hoisting",
    "06 10 00 Rough Carpentry",
    "06 20 00 Interior Trim / Finish Carpentry",
    "06 40 00 Architectural Woodwork",
    "08 11 00 Hollow Metal Frames & Doors",
    "08 14 00 Wood Doors",
    "08 17 13 Exterior Pre-Hung Doors",
    "08 17 43 Interior Pre-Hung Doors",
    "08 50 00 Windows",
    "08 71 00 Door Hardware",
    "08 83 00 Mirrors",
    "10 21 13 Toilet Compartments",
    "10 28 00 Bath Accessories",
    "10 44 00 Fire Specialties",
    "10 55 00 Postal Specialties",
    "10 56 23 Wire Shelving",
    "12 21 00 Window Blinds",
]
NSEC = 25  # rows reserved for sections (extra blanks for custom ones)
UOMS = ["LF", "EA", "SF", "SET", "PR", "LS", "DAY", "MO", "BOX", "HR", "SHT"]
PTYPES = ["Multifamily", "Senior Living", "Hospitality", "Healthcare", "Education", "Office / Retail", "Other"]

FIRST, LAST = 6, 505  # Line Detail data rows (500 lines)
LD = "'Line Detail'"
def rng(col):
    return f"{LD}!${col}${FIRST}:${col}${LAST}"

# ---------------------------------------------------------------- Instructions
ws0 = wb.active
ws0.title = "Instructions"

# ---------------------------------------------------------------- Job Info
ji = wb.create_sheet("Job Info")
ji.column_dimensions["A"].width = 38
ji.column_dimensions["B"].width = 42
ji.column_dimensions["C"].width = 60
ji["A1"] = "JOB INFO & PRICING INPUTS"; ji["A1"].font = f_title
ji["A2"] = "Yellow cells with blue text are inputs. Everything else calculates."; ji["A2"].font = Font(name=F, size=9, italic=True)

J = {}  # key -> cell address
rows = [
    ("COMPANY", None, None, None),
    ("co_name", "Company Name", "Four Corners Building Supply", None),
    ("co_addr", "Address", "4349 Corporate Road", None),
    ("co_city", "City / State / Zip", "North Charleston, SC 29405", None),
    ("co_phone", "Office Phone", "(843) 970-3812", None),
    ("co_web", "Website", "www.buildfcbs.com", None),
    (None, None, None, None),
    ("CUSTOMER / GENERAL CONTRACTOR", None, None, None),
    ("gc", "General Contractor", None, None),
    ("attn", "Attention", None, None),
    ("gc_addr", "Address", None, None),
    ("gc_city", "City / State / Zip", None, None),
    ("gc_phone", "Phone", None, None),
    ("gc_email", "Email", None, None),
    (None, None, None, None),
    ("PROJECT", None, None, None),
    ("job", "Job Name", None, None),
    ("jobno", "Four Corners Job #", None, None),
    ("job_addr", "Job Address", None, None),
    ("job_city", "Job City / State / Zip", None, None),
    ("ptype", "Project Type", "Multifamily", None),
    ("plans", "Plans / Specs Dated", None, "mm/dd/yyyy"),
    ("addenda", "Addenda Acknowledged", "None", None),
    ("bid_date", "Proposal Date", None, "mm/dd/yyyy"),
    ("prop_no", "Proposal #", None, None),
    ("rev", "Revision", "0", None),
    (None, None, None, None),
    ("ESTIMATOR", None, None, None),
    ("est", "Estimator Name", None, None),
    ("est_phone", "Estimator Phone", None, None),
    ("est_email", "Estimator Email", None, None),
    (None, None, None, None),
    ("PRICING & COMMERCIAL TERMS", None, None, None),
    ("tax", "Sales Tax Rate on Material", 0.09, "0.00%"),
    ("tax_on", "Include Tax in Material Prices?", "Yes", None),
    ("mmarg", "Default Material Margin (GP %)", 0.20, "0.0%"),
    ("lmarg", "Default Labor Margin (GP %)", 0.15, "0.0%"),
    ("bond_on", "Payment & Performance Bond Required?", "No", None),
    ("bond_rate", "Bond Rate (% of base bid)", 0.015, "0.00%"),
    ("valid", "Proposal Valid (days)", 30, "0"),
    ("esc", "Escalation Threshold Before Pass-Through", 0.03, "0.0%"),
    ("award", "Contract & Approved Submittals By", None, "mm/dd/yyyy"),
    ("compl", "Substantial Completion Date", None, "mm/dd/yyyy"),
    ("terms", "Payment Terms", "Monthly progress billing, Net 30 from invoice date", None),
    ("ret", "Max Retainage", 0.05, "0.0%"),
    ("mob", "Mobilizations Included (per bldg/phase, per scope)", 2, "0"),
    ("mob_rate", "Additional Mobilization Charge ($/trip)", 350, CUR0),
    (None, None, None, None),
    ("AREAS (used for the Area Breakout - rename to fit the job)", None, None, None),
    ("a1", "Area 1", "Units / Apartments", None),
    ("a2", "Area 2", "Clubhouse", None),
    ("a3", "Area 3", "Leasing / Amenity", None),
    ("a4", "Area 4", "Corridors / Common", None),
    ("a5", "Area 5", "Maintenance", None),
    ("a6", "Area 6", "Garages", None),
    ("a7", "Area 7", "Pool Cabana", None),
    ("a8", "Area 8", None, None),
]
notes = {
    "tax": "9% = Charleston/Berkeley Co. (Carter used 9% on Goose Creek, 7% on Bluffton). VERIFY the county rate for every job.",
    "mmarg": "Gross margin: sell = cost / (1 - margin). Can be overridden per line on Line Detail.",
    "lmarg": "Gross margin on installer cost. Can be overridden per line on Line Detail.",
    "bond_rate": "1.5% matches the bond Carter carried on Pepper Hall ($36,008 on ~$2.40M). Get a real quote from the surety on bonded jobs.",
    "valid": "Carter used 30 days.",
    "esc": "Material increases above this % get passed to the GC (Qualification #3).",
    "ret": "Shows up in the payment qualification.",
    "mob": "Carter: 'reasonable number of units per mobilization'. We name a number instead.",
    "mob_rate": "Placeholder; set your own trip charge.",
    "prop_no": "Suggest: FC-YYMM-### (e.g. FC-2610-001).",
    "plans": "Leave blank to print '[date TBD]'.",
}
r = 4
for key, label, val, fmt in rows:
    if key is None:
        r += 1; continue
    if label is None:
        c = ji.cell(row=r, column=1, value=key); c.font = f_sec
        ji.cell(row=r, column=1).border = Border(bottom=Side(style="medium", color=NAVY))
        ji.cell(row=r, column=2).border = Border(bottom=Side(style="medium", color=NAVY))
        r += 1; continue
    ji.cell(row=r, column=1, value=label).font = f_norm
    c = ji.cell(row=r, column=2)
    inp(c, val, fmt)
    if key in notes:
        n = ji.cell(row=r, column=3, value=notes[key]); n.font = Font(name=F, size=9, italic=True, color="595959")
    J[key] = f"'Job Info'!$B${r}"
    J[key + "_r"] = r
    r += 1

# margin check (internal)
r += 1
ji.cell(row=r, column=1, value="JOB MARGIN CHECK (INTERNAL - DO NOT SEND)").font = f_sec
ji.cell(row=r, column=1).border = Border(bottom=Side(style="medium", color=NAVY))
ji.cell(row=r, column=2).border = Border(bottom=Side(style="medium", color=NAVY))
r += 1
chk = [
    ("Base Material Cost (pre-tax)", f'=SUMIFS({rng("U")},{rng("B")},"Base")', CUR0),
    ("Base Labor Cost", f'=SUMIFS({rng("V")},{rng("B")},"Base")', CUR0),
    ("Sales Tax in Base Bid", f'=SUMIFS({rng("X")},{rng("B")},"Base")', CUR0),
    ("Base Sell (before bond)", f'=SUMIFS({rng("R")},{rng("B")},"Base")', CUR0),
    ("Gross Profit $ (excl. tax)", "=B{s}-B{m}-B{l}-B{t}", CUR0),
    ("Gross Profit % (excl. tax)", "=IF(B{s}-B{t}=0,0,B{g}/(B{s}-B{t}))", PCT),
    ("Labor % of Base Sell", "=IF(B{s}=0,0,SUMIFS(" + rng("Q") + "," + rng("B") + ',"Base")/B{s})', PCT),
    ("Base lines with no CSI section ($)", f'=SUMIFS({rng("R")},{rng("B")},"Base",{rng("C")},"")', CUR0),
    ("Lines with a description but no Bid Group (count)", f'=COUNTIFS({rng("F")},"?*",{rng("B")},"")', "0"),
    ("CSI sections carrying $ (Proposal fits 16)", "=Lists!$J$2", "0"),
]
base = r
m, l, t, s, g = base, base + 1, base + 2, base + 3, base + 4
for i, (lab, fml, fmt) in enumerate(chk):
    ji.cell(row=r, column=1, value=lab).font = f_norm
    c = ji.cell(row=r, column=2, value=fml.format(m=m, l=l, t=t, s=s, g=g))
    c.number_format = fmt; c.font = f_bold if i in (4, 5) else f_link; c.border = box
    r += 1
ji.cell(row=base + 6, column=3, value="Carter's bids ran 24-29% labor. Useful sanity check.").font = Font(name=F, size=9, italic=True, color="595959")
ji.cell(row=base + 7, column=3, value="Should be $0. Anything here is missing from the proposal section list.").font = Font(name=F, size=9, italic=True, color="C00000")
ji.cell(row=base + 8, column=3, value="Should be 0. Lines without a Bid Group are NOT priced anywhere.").font = Font(name=F, size=9, italic=True, color="C00000")
ji.cell(row=base + 9, column=3, value="If over 16, combine sections - extras won't show on the Proposal summary.").font = Font(name=F, size=9, italic=True, color="595959")
ji.freeze_panes = "A4"

AREA_CELLS = [J[f"a{i}"] for i in range(1, 9)]

# ---------------------------------------------------------------- Lists sheet
ls = wb.create_sheet("Lists")
for col, (h, w) in enumerate([("Bid Group", 12), ("CSI Section", 44), ("Base $", 14), ("Seq", 6),
                               ("UOM", 8), ("Yes/No", 8), ("Project Type", 18), ("Areas", 24)], 1):
    ls.cell(row=1, column=col, value=h)
    ls.column_dimensions[L(col)].width = w
hdr(ls, 1, ["Bid Group", "CSI Section", "Base $", "Seq", "UOM", "Yes/No", "Project Type", "Areas"])
for i, g_ in enumerate(GROUPS): ls.cell(row=2 + i, column=1, value=g_)
for i in range(NSEC):
    rr = 2 + i
    c = ls.cell(row=rr, column=2, value=SECTIONS[i] if i < len(SECTIONS) else None)
    if i >= len(SECTIONS):
        inp(c)
    ls.cell(row=rr, column=3, value=f'=IF(B{rr}="",0,SUMIFS({rng("R")},{rng("C")},B{rr},{rng("B")},"Base"))').number_format = CUR0
    ls.cell(row=rr, column=4, value=f'=IF(C{rr}=0,"",SUMPRODUCT(--($C$2:C{rr}<>0)))')
for i, u in enumerate(UOMS): ls.cell(row=2 + i, column=5, value=u)
ls.cell(row=2, column=6, value="Yes"); ls.cell(row=3, column=6, value="No")
for i, p in enumerate(PTYPES): ls.cell(row=2 + i, column=7, value=p)
for i, a in enumerate(AREA_CELLS): ls.cell(row=2 + i, column=8, value=f'=IF({a}="","",{a})')
ls.cell(row=NSEC + 3, column=2, value="Blank yellow rows above = add custom CSI sections. Keep the 'NN NN NN Name' format.").font = Font(name=F, size=9, italic=True)
SEC_RNG = f"Lists!$B$2:$B${NSEC + 1}"

def dv_list(ws, formula, ref, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if prompt:
        dv.promptTitle, dv.prompt = "Pick from list", prompt
    ws.add_data_validation(dv); dv.add(ref)

yn = "=Lists!$F$2:$F$3"
for k in ("tax_on", "bond_on"):
    dv_list(ji, yn, f"B{J[k + '_r']}")
dv_list(ji, f"=Lists!$G$2:$G${1 + len(PTYPES)}", f"B{J['ptype_r']}")

# ---------------------------------------------------------------- Line Detail
ld = wb.create_sheet("Line Detail")
ld["A1"] = "LINE DETAIL / TAKEOFF (INTERNAL - contains cost)"; ld["A1"].font = f_title
ld["A2"] = "=\"Job: \"&IF(" + J["job"] + "=\"\",\"[Job Name]\"," + J["job"] + ")&\"   |   Tax on material: \"&IF(" + J["tax_on"] + "=\"Yes\",TEXT(" + J["tax"] + ",\"0.00%\"),\"not included\")&\"   |   Default margins: Mat \"&TEXT(" + J["mmarg"] + ",\"0.0%\")&\" / Labor \"&TEXT(" + J["lmarg"] + ",\"0.0%\")"
ld["A2"].font = f_bold
ld["A3"] = "Enter lines grouped by CSI section, then by area, so the Bid Summary prints in a sensible order. Put each hardware set's components in the Description (one line per set, like Carter does)."
ld["A3"].font = Font(name=F, size=9, italic=True)
cols = [
    ("Line #", 7), ("Bid Group", 10), ("CSI Section", 34), ("Area", 18), ("Mark / Tag", 12),
    ("Description / Specification", 60), ("Qty", 11), ("UOM", 7), ("Mat Unit Cost", 12),
    ("Labor Unit Cost", 12), ("Mat Margin Override", 10), ("Labor Margin Override", 10),
    ("Mat Unit Sell (incl. tax)", 12), ("Labor Unit Sell", 12), ("Unit Price (Mat + Labor)", 12),
    ("Material Ext.", 14), ("Labor Ext.", 14), ("Total Ext.", 14),
    ("Customer Note (prints)", 28), ("Internal Note (does not print)", 28),
    ("Mat Cost Ext.", 13), ("Labor Cost Ext.", 13), ("Print Seq", 8), ("Tax $", 11),
]
for i, (h, w) in enumerate(cols, 1):
    ld.column_dimensions[L(i)].width = w
hdr(ld, 5, [c[0] for c in cols])
ld.row_dimensions[5].height = 42
# group header band
for rng_, lab in (("A4:H4", "WHAT"), ("I4:L4", "COST INPUTS"), ("M4:R4", "SELL (calculated)"), ("S4:T4", "NOTES"), ("U4:X4", "HELPERS")):
    ld.merge_cells(rng_)
    c = ld[rng_.split(":")[0]]; c.value = lab; c.font = f_bold; c.alignment = center; c.fill = fill_band

tax_mult = f'(1+IF({J["tax_on"]}="Yes",{J["tax"]},0))'
for rr in range(FIRST, LAST + 1):
    ld.cell(row=rr, column=1, value=f'=IF(F{rr}="","",ROW()-{FIRST - 1})').font = f_norm
    for col in range(2, 13):
        inp(ld.cell(row=rr, column=col))
    ld.cell(row=rr, column=7).number_format = QTY
    for col in (9, 10):
        ld.cell(row=rr, column=col).number_format = CUR
    for col in (11, 12):
        ld.cell(row=rr, column=col).number_format = PCT
    ld.cell(row=rr, column=6).alignment = wrap
    mm = f'IF(K{rr}="",{J["mmarg"]},K{rr})'
    lm = f'IF(L{rr}="",{J["lmarg"]},L{rr})'
    ld.cell(row=rr, column=13, value=f'=IF(F{rr}="","",N(I{rr})/(1-{mm})*{tax_mult})')
    ld.cell(row=rr, column=14, value=f'=IF(F{rr}="","",N(J{rr})/(1-{lm}))')
    ld.cell(row=rr, column=15, value=f'=IF(F{rr}="","",M{rr}+N{rr})')
    ld.cell(row=rr, column=16, value=f'=IF(F{rr}="","",N(G{rr})*M{rr})')
    ld.cell(row=rr, column=17, value=f'=IF(F{rr}="","",N(G{rr})*N{rr})')
    ld.cell(row=rr, column=18, value=f'=IF(F{rr}="","",P{rr}+Q{rr})')
    for col in (19, 20):
        inp(ld.cell(row=rr, column=col)); ld.cell(row=rr, column=col).alignment = wrap
    ld.cell(row=rr, column=21, value=f'=IF(F{rr}="","",N(G{rr})*N(I{rr}))')
    ld.cell(row=rr, column=22, value=f'=IF(F{rr}="","",N(G{rr})*N(J{rr}))')
    ld.cell(row=rr, column=23, value=f'=IF(F{rr}="","",IF(B{rr}="Excluded","",COUNTIFS($F${FIRST}:F{rr},"?*",$B${FIRST}:B{rr},"<>Excluded")))')
    ld.cell(row=rr, column=24, value=f'=IF(F{rr}="","",N(G{rr})*N(I{rr})/(1-{mm})*IF({J["tax_on"]}="Yes",{J["tax"]},0))')
    for col in range(13, 19):
        ld.cell(row=rr, column=col).number_format = CUR
        ld.cell(row=rr, column=col).font = f_bold if col == 18 else f_norm
    for col in (21, 22, 24):
        ld.cell(row=rr, column=col).number_format = CUR
        ld.cell(row=rr, column=col).font = f_norm
    ld.cell(row=rr, column=23).font = f_norm

# totals row above header? put totals at row 4? Use a totals row below
tr = LAST + 1
ld.cell(row=tr, column=6, value="TOTAL - BASE BID LINES").font = f_bold
for col in (16, 17, 18, 21, 22):
    cl = L(col)
    c = ld.cell(row=tr, column=col, value=f'=SUMIFS({cl}{FIRST}:{cl}{LAST},$B${FIRST}:$B${LAST},"Base")')
    c.number_format = CUR0; c.font = f_bold; c.fill = fill_tot; c.border = box

dv_list(ld, f"=Lists!$A$2:$A${1 + len(GROUPS)}", f"B{FIRST}:B{LAST}", "Base, Alt 1-5, Optional, or Excluded (excluded lines don't print)")
dv_list(ld, f"={SEC_RNG}", f"C{FIRST}:C{LAST}")
dv_list(ld, "=Lists!$H$2:$H$9", f"D{FIRST}:D{LAST}", "Areas come from Job Info")
dv_list(ld, f"=Lists!$E$2:$E${1 + len(UOMS)}", f"H{FIRST}:H{LAST}")
red = PatternFill("solid", fgColor="F8CBAD")
ld.conditional_formatting.add(f"B{FIRST}:B{LAST}", FormulaRule(formula=[f'AND($F{FIRST}<>"",$B{FIRST}="")'], fill=red))
ld.conditional_formatting.add(f"C{FIRST}:C{LAST}", FormulaRule(formula=[f'AND($F{FIRST}<>"",$C{FIRST}="")'], fill=red))
ld.conditional_formatting.add(f"G{FIRST}:G{LAST}", FormulaRule(formula=[f'AND($F{FIRST}<>"",$G{FIRST}="")'], fill=red))
ld.freeze_panes = f"G{FIRST}"
ld.auto_filter.ref = f"A5:X{LAST}"

# Example rows (realistic, drawn from the Carter bids' scope)
examples = [
    ("Base", "06 20 00 Interior Trim / Finish Carpentry", "Units / Apartments", "", "Base FJ Primed 1x6", 1000, "LF", 1.45, 1.10, None, None, "", "EXAMPLE - delete before use"),
    ("Base", "08 17 43 Interior Pre-Hung Doors", "Units / Apartments", "Type 4", '3068 Hollow Core Hardboard 2-Panel Shaker 1-3/8", FJ split jamb 4-7/8" w/ FJ 1x4 casing, 3 ea US15 hinges, 2-3/8" backset single bore', 10, "EA", 165.00, 38.00, None, None, "", "EXAMPLE - delete before use"),
    ("Base", "08 71 00 Door Hardware", "Units / Apartments", "HW Set Bed/Bath", "1 ea privacy lever US15; 1 ea door stop US15", 10, "SET", 18.50, 9.50, None, None, "", "EXAMPLE - delete before use"),
    ("Alt 1", "08 71 00 Door Hardware", "Units / Apartments", "HW Set Entry", "ADD: battery-powered keypad deadbolt in lieu of keyed deadbolt", 10, "EA", 240.00, 25.00, None, None, "Add alternate", "EXAMPLE - delete before use"),
]
for i, ex in enumerate(examples):
    rr = FIRST + i
    for j, v in enumerate(ex[:11]):
        if v is not None:
            c = ld.cell(row=rr, column=2 + j, value=v); c.font = f_ex
    ld.cell(row=rr, column=19, value=ex[11] or None).font = f_ex
    ld.cell(row=rr, column=20, value=ex[12]).font = f_ex

# ---------------------------------------------------------------- Scope Terms
st = wb.create_sheet("Scope Terms")
st.column_dimensions["A"].width = 15
st.column_dimensions["B"].width = 10
st.column_dimensions["C"].width = 110
st.column_dimensions["D"].width = 6
st.column_dimensions["E"].width = 16
st["A1"] = "INCLUSIONS / EXCLUSIONS / QUALIFICATIONS / TERMS"; st["A1"].font = f_title
st["A2"] = "Set column B to Yes/No per job. 'Yes' items print on the Proposal, auto-numbered. Edit wording freely; add your own items in the blank rows at the end of each group (set the Category)."
st["A2"].font = Font(name=F, size=9, italic=True)
st["A3"] = "Items in green are formulas that pull from Job Info - overwrite them only if you need custom wording."
st["A3"].font = Font(name=F, size=9, italic=True, color="008000")
hdr(st, 4, ["Category", "Print?", "Text", "Seq", "Key"])

def dt(key, tbd="[date TBD]"):
    return f'IF({J[key]}="","{tbd}",TEXT({J[key]},"mmmm d, yyyy"))'

INC = [
    "All items listed in the attached Bid Summary, which is part of this proposal.",
    f'=IF({J["tax_on"]}="Yes","All material pricing includes "&TEXT({J["tax"]},"0.0%")&" sales tax.","Sales tax is NOT included in material pricing and will be added at invoicing.")',
    f'="Pricing is based on plans and specifications dated "&{dt("plans")}&", and addenda: "&{J["addenda"]}&"."',
    "Delivery of all materials to the jobsite, ground level, at each building.",
    "Installation labor for every line item that shows a labor price in the Bid Summary.",
    "Interior and exterior pre-hung doors with hinges installed and bored or prepped for the scheduled hardware.",
    "Fire-rated doors and frames with factory labels as scheduled.",
    "Product data, door and hardware schedules, and submittals; warranty and O&M documents at closeout.",
    "Removal of our packaging and debris to the GC-provided on-site dumpster.",
]
EXC = [
    ("Yes", "Chair rail and crown moulding in units"),
    ("Yes", "Cased openings and interior columns"),
    ("Yes", "Cabinets, countertops, vanities, desks, and built-ins unless listed"),
    ("Yes", "Stair parts and railings"),
    ("Yes", "Toilet partitions, shower doors, and shower enclosures unless listed"),
    ("Yes", "Garage doors, overhead doors, and draft-stop or attic access doors"),
    ("Yes", "Aluminum storefront, curtain wall, and associated door hardware"),
    ("Yes", "Glass and glazing for hollow metal doors and frames unless listed"),
    ("Yes", "Flashing, sealants, and weatherproofing for windows and doors"),
    ("Yes", "Painting, staining, caulking, and nail-hole fill (trim and doors are primed or prefinished as noted)"),
    ("Yes", "Wood blocking, backing, and framing"),
    ("Yes", "Grouting of hollow metal frames (material and labor)"),
    ("Yes", "Electrified hardware, access control, power supplies, wiring, and low-voltage work (Division 28)"),
    ("Yes", "Final keying, permanent cores, and master key systems unless listed"),
    ("Yes", "Grab bars installed in tile or solid surface (price on request)"),
    ("Yes", "Reclaimed or specialty-species wood trim unless listed"),
    ("Yes", "Hoisting and stocking above ground level unless listed under 01 54 16"),
    ("Yes", "Off-site trash removal and dumpsters"),
    ("Yes", "Composite clean-up participation and clean-up back charges"),
    ("Yes", "OCIP / CCIP insurance programs and related costs"),
    ("Yes", '=IF(' + J["bond_on"] + '="Yes","Bonds other than the payment and performance bond shown in the Proposal Summary","Payment and performance bonds")'),
    ("Yes", "Builder's risk insurance, permits, inspection fees, and testing"),
    ("Yes", "Liquidated and consequential damages"),
    ("Yes", "Prevailing wage / Davis-Bacon labor rates unless stated"),
    ("Yes", "Attic stock and overage material"),
    ("Yes", "Protection of installed work from damage by other trades"),
    ("Yes", "Engineering, delegated design, and sealed shop drawings"),
    ("No", "Window blinds and window treatments"),
    ("No", "Postal boxes and mail kiosks"),
    ("No", "Wire shelving"),
    ("No", "Bath accessories"),
    ("No", "Fire extinguishers and cabinets"),
]
QUAL = [
    f'="This proposal is valid for "&{J["valid"]}&" days from the proposal date."',
    f'="Pricing is based on an executed contract and approved submittals by "&{dt("award")}&", and substantial completion by "&{dt("compl")}&". If the schedule moves past these dates, Four Corners may revise pricing."',
    f'="Material pricing reflects current market conditions. If a material cost goes up more than "&TEXT({J["esc"]},"0%")&" from proposal pricing before procurement, Four Corners will submit the increase to the GC with supporting documentation before ordering."',
    "Lead times come from current manufacturer quotes and are confirmed at release. Special-order items need approved submittals and a written release before ordering, and can't be cancelled or returned once released.",
    f'="Payment terms: "&{J["terms"]}&". Retainage may not exceed "&TEXT({J["ret"]},"0%")&" and is released at substantial completion of our scope."',
    "GC provides hoisting and material stocking for buildings over four stories, unless hoisting is listed in this proposal.",
    "GC provides all-weather access and temporary roads to each building for deliveries.",
    "GC provides secure, dry on-site storage. If on-site storage isn't available, storage charges will apply.",
    "Work follows a mutually agreed construction schedule. Areas must be dried in, and conditioned where specified, before interior trim and doors are installed.",
    f'="Pricing includes "&{J["mob"]}&" mobilization(s) per building or phase for each scope. Extra mobilizations caused by others will be billed at "&TEXT({J["mob_rate"]},"$#,##0")&" per trip."',
    "Supervision is provided as needed; it's not full-time on site unless stated.",
    "This is a lump-sum proposal, broken out by section for the GC's convenience. If less than the full scope is awarded, Four Corners may revise pricing.",
    "Quantities are our takeoff from the documents provided and aren't guaranteed. Please review the attached Bid Summary and let us know right away about any discrepancy. Changes in quantity, scope, or product will be priced as a change order.",
    "Where the drawings and specifications conflict, pricing is based on the product described in the attached Bid Summary. Please flag any conflict before award.",
    "Fees for third-party payment management or compliance systems (such as Textura or GCPay) are not included.",
    "Changes require a written change order signed by both parties before the work proceeds.",
    "Back charges won't be accepted without prior written notice and an opportunity to correct.",
    "Repairs to installed work damaged by other trades will be billed as extra work.",
    "Manufacturer warranties apply to all products. Four Corners warrants its installation workmanship for one (1) year from substantial completion.",
    "Four Corners carries standard insurance and will provide a certificate of insurance. Additional insured and waiver of subrogation are provided per a mutually agreed subcontract.",
]
TERMS = [
    "This proposal assumes Four Corners and the Customer can agree on contract terms. If the subcontract contains terms Four Corners can't accept, this proposal is null and void.",
    "All work is subject to credit approval. Four Corners may decline any project for credit reasons.",
    "Customer pays according to the Four Corners credit agreement, unless Four Corners accepts alternate terms in writing.",
    "Past-due balances accrue a finance charge of 1.5% per month. Customer is responsible for collection costs and reasonable attorney's fees.",
    "Customer is responsible for all permits and inspections required by local agencies.",
    "If work can't start within a reasonable time of the projected start date, Four Corners may revise this proposal.",
    "Four Corners keeps all lien rights under South Carolina law.",
]
groups = [("Inclusion", [("Yes", t) for t in INC]), ("Exclusion", EXC),
          ("Qualification", [("Yes", t) for t in QUAL]), ("Term", [("Yes", t) for t in TERMS])]
r = 5
TEXT_LEN = {}  # category -> list of default text lengths
for cat, items in groups:
    for yn_, text in items + [("No", None)] * 4:
        st.cell(row=r, column=1, value=cat).font = f_norm
        inp(st.cell(row=r, column=2), yn_)
        c = st.cell(row=r, column=3, value=text)
        c.alignment = wrap
        if text and str(text).startswith("="):
            c.font = f_link; c.border = box
        else:
            inp(c, text)
            c.alignment = wrap
        st.cell(row=r, column=4, value=f'=IF(AND(B{r}="Yes",LEN(C{r})>0),COUNTIFS($A$5:A{r},A{r},$B$5:B{r},"Yes",$C$5:C{r},"?*"),"")').font = f_norm
        # COUNTIFS "?*" ignores formula text? it counts any text value - fine.
        st.cell(row=r, column=5, value=f'=IF(D{r}="","",A{r}&"|"&D{r})').font = f_norm
        if text:
            t_ = str(text)
            est_ = len(t_) if not t_.startswith("=") else (max([len(x) for x in re.findall(r'"([^"]*)"', t_)] or [60]) if t_.startswith("=IF(") else sum(len(x) for x in re.findall(r'"([^"]*)"', t_)) + 12)
            TEXT_LEN.setdefault(cat, []).append(est_)
        r += 1
ST_LAST = r - 1
dv_list(st, yn, f"B5:B{ST_LAST}")
dv_list(st, '"Inclusion,Exclusion,Qualification,Term"', f"A5:A{ST_LAST}")
st.freeze_panes = "A5"

# ---------------------------------------------------------------- Alternates
al = wb.create_sheet("Alternates")
al.column_dimensions["A"].width = 10
al.column_dimensions["B"].width = 70
al.column_dimensions["C"].width = 16
al.column_dimensions["D"].width = 16
al.column_dimensions["E"].width = 16
al["A1"] = "ALTERNATES & OPTIONS"; al["A1"].font = f_title
al["A2"] = "Alternate amounts total automatically from Line Detail lines tagged Alt 1-Alt 5. Enter deducts as negative quantities. Describe each alternate in column B; it only prints if it has a description or an amount."
al["A2"].font = Font(name=F, size=9, italic=True)
hdr(al, 4, ["Alt #", "Description (prints on proposal)", "Add / (Deduct)", "Material", "Labor"])
for i in range(5):
    rr = 5 + i
    g_ = f"Alt {i + 1}"
    al.cell(row=rr, column=1, value=g_).font = f_bold
    inp(al.cell(row=rr, column=2))
    for col, src in ((3, "R"), (4, "P"), (5, "Q")):
        c = al.cell(row=rr, column=col, value=f'=SUMIFS({rng(src)},{rng("B")},A{rr})')
        c.number_format = CUR0; c.font = f_link; c.border = box
al["B5"].value = "Battery-powered keypad deadbolts at unit entries in lieu of keyed deadbolts (EXAMPLE)"
al["B5"].font = f_ex
al.cell(row=11, column=1, value="Optional").font = f_bold
inp(al.cell(row=11, column=2), "Optional items, priced individually in the attached Bid Summary")
for col, src in ((3, "R"), (4, "P"), (5, "Q")):
    c = al.cell(row=11, column=col, value=f'=SUMIFS({rng(src)},{rng("B")},"Optional")')
    c.number_format = CUR0; c.font = f_link; c.border = box

# ---------------------------------------------------------------- Proposal (customer-facing)
pr = wb.create_sheet("Proposal")
widths = {"A": 6, "B": 16, "C": 30, "D": 14, "E": 16, "F": 18}
for k, v in widths.items():
    pr.column_dimensions[k].width = v
TEXTW = 124  # approx chars per line across B:F

def merge_val(ws, rngs, val, font=f_norm, align=None, fill=None, fmt=None, border=None):
    ws.merge_cells(rngs)
    c = ws[rngs.split(":")[0]]
    c.value = val; c.font = font
    if align: c.alignment = align
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if border: c.border = border
    return c

merge_val(pr, "A1:F1", f'={J["co_name"]}', Font(name=F, size=18, bold=True, color=NAVY), Alignment(horizontal="left"))
merge_val(pr, "A2:F2", f'={J["co_addr"]}&"  |  "&{J["co_city"]}&"  |  "&{J["co_phone"]}&"  |  "&{J["co_web"]}', Font(name=F, size=9, color="595959"))
for col in range(1, 7):
    pr.cell(row=3, column=col).border = Border(bottom=Side(style="thick", color=NAVY))
merge_val(pr, "A5:F5", "PROPOSAL", Font(name=F, size=16, bold=True, color=NAVY), Alignment(horizontal="center"))

def kv(row, lab, val, col_l="A", rng_v="B{r}:C{r}", lab2=None, val2=None, fmt2=None):
    c = pr[f"{col_l}{row}"]; c.value = lab; c.font = f_bold
    merge_val(pr, rng_v.format(r=row), val)
    if lab2:
        pr[f"D{row}"].value = lab2; pr[f"D{row}"].font = f_bold
        c2 = merge_val(pr, f"E{row}:F{row}", val2, align=Alignment(horizontal="left"))
        if fmt2: c2.number_format = fmt2

pr.merge_cells("A7:A7")
kv(7, "To:", f'={J["gc"]}&""', lab2="Date:", val2=f'=IF({J["bid_date"]}="","",{J["bid_date"]})', fmt2="mmmm d, yyyy")
kv(8, "Attn:", f'={J["attn"]}&""', lab2="Proposal #:", val2=f'={J["prop_no"]}&IF({J["rev"]}&""="0",""," Rev "&{J["rev"]})')
kv(9, "Address:", f'={J["gc_addr"]}&""', lab2="From:", val2=f'={J["est"]}&""')
kv(10, "", f'={J["gc_city"]}&""', lab2="Phone:", val2=f'={J["est_phone"]}&""')
kv(11, "Phone:", f'={J["gc_phone"]}&""', lab2="Email:", val2=f'={J["est_email"]}&""')
for rr in range(7, 12):
    pr[f"A{rr}"].alignment = Alignment(horizontal="left")

hdr(pr, 13, ["", "JOB NAME", "", "JOB #", "JOB ADDRESS", ""])
pr.merge_cells("A13:C13"); pr["A13"].value = "JOB NAME"
pr.merge_cells("E13:F13")
merge_val(pr, "A14:C14", f'={J["job"]}&""', f_bold, center)
pr["D14"].value = f'={J["jobno"]}&""'; pr["D14"].alignment = center; pr["D14"].font = f_norm
merge_val(pr, "E14:F14", f'={J["job_addr"]}&IF({J["job_city"]}="","",", "&{J["job_city"]})', f_norm, center)
pr.row_dimensions[14].height = 28

r = 16
hdr(pr, r, ["", "PROPOSAL SUMMARY", "", "", "", "AMOUNT"])
pr.merge_cells(f"A{r}:E{r}"); pr[f"A{r}"].value = "PROPOSAL SUMMARY (BY CSI SECTION)"
r += 1
# helper calcs on Lists (fixed cells other sheets can reference)
ls["J1"].value = "Proposal calcs"; ls["J1"].font = f_bold
ls["J2"].value = f"=MAX($D$2:$D${NSEC + 1})"; ls["I2"].value = "# sections w/ $"
ls["J3"].value = f'=SUMIFS({rng("R")},{rng("B")},"Base")'; ls["I3"].value = "Subtotal"
ls["J4"].value = f'=IF({J["bond_on"]}="Yes",ROUND(J3*{J["bond_rate"]},2),0)'; ls["I4"].value = "Bond"
ls["J5"].value = "=J3+J4"; ls["I5"].value = "Base bid total"
ls["J6"].value = f'=IF({J["bond_on"]}="Yes","Payment & Performance Bond ("&TEXT({J["bond_rate"]},"0.0%")&")","Payment & Performance Bond - Not Included")'; ls["I6"].value = "Bond label"
for rr_ in range(3, 6): ls[f"J{rr_}"].number_format = CUR
ls.column_dimensions["I"].width = 16; ls.column_dimensions["J"].width = 40
BOND_REF = "Lists!$J$4"

sum_first = r
NSUM = 16
n_ = "Lists!$J$2"
for k in range(1, NSUM + 4):
    sec = f'INDEX(Lists!$B$2:$B${NSEC + 1},MATCH({k},Lists!$D$2:$D${NSEC + 1},0))'
    amt = f'INDEX(Lists!$C$2:$C${NSEC + 1},MATCH({k},Lists!$D$2:$D${NSEC + 1},0))'
    lab = f'=IF({k}<={n_},{sec},IF({k}={n_}+1,"Subtotal",IF({k}={n_}+2,Lists!$J$6,IF({k}={n_}+3,"BASE BID TOTAL",""))))'
    val = f'=IF({k}<={n_},{amt},IF({k}={n_}+1,Lists!$J$3,IF({k}={n_}+2,Lists!$J$4,IF({k}={n_}+3,Lists!$J$5,""))))'
    merge_val(pr, f"B{r}:E{r}", lab)
    c = pr.cell(row=r, column=6, value=val); c.number_format = CUR; c.font = f_norm
    r += 1
sum_last = r - 1
fill_total = PatternFill("solid", fgColor=NAVY, bgColor=NAVY)
pr.conditional_formatting.add(f"A{sum_first}:F{sum_last}", FormulaRule(formula=[f'$B{sum_first}="BASE BID TOTAL"'], fill=fill_total, font=Font(name=F, bold=True, color="FFFFFF")))
pr.conditional_formatting.add(f"A{sum_first}:F{sum_last}", FormulaRule(formula=[f'$B{sum_first}="Subtotal"'], font=Font(name=F, bold=True), border=Border(top=thin)))
r += 2

hdr(pr, r, ["", "ALTERNATES & OPTIONS (NOT INCLUDED IN BASE BID)", "", "", "", "ADD / (DEDUCT)"])
pr.merge_cells(f"A{r}:E{r}"); pr[f"A{r}"].value = "ALTERNATES & OPTIONS (NOT INCLUDED IN BASE BID)"
r += 1
alt_first = r
for i in range(6):
    ar = 5 + i if i < 5 else 11
    show = f'OR(Alternates!$B${ar}<>"",Alternates!$C${ar}<>0)' if i < 5 else f'Alternates!$C${ar}<>0'
    pr.cell(row=r, column=1, value=f'=IF({show},Alternates!$A${ar},"")').font = f_bold
    merge_val(pr, f"B{r}:E{r}", f'=IF({show},Alternates!$B${ar}&"","")', align=wrap)
    c = pr.cell(row=r, column=6, value=f'=IF({show},Alternates!$C${ar},"")'); c.number_format = CUR
    pr.row_dimensions[r].height = 15
    r += 1
r += 1

intro = (f'={J["co_name"]}&" is pleased to present this proposal for "&IF({J["job"]}="","[Job Name]",{J["job"]})&'
         f'", based on plans and specifications dated "&{dt("plans")}&", subject to the inclusions, exclusions, qualifications, and terms below. The attached Bid Summary is part of this proposal."')
merge_val(pr, f"A{r}:F{r}", intro, align=wrap)
pr.row_dimensions[r].height = 42
r += 2

def block(title, cat, n, est_lens):
    global r
    hdr(pr, r, ["", title, "", "", "", ""])
    pr.merge_cells(f"A{r}:F{r}"); pr[f"A{r}"].value = title
    pr[f"A{r}"].alignment = Alignment(horizontal="left", vertical="center")
    r += 1
    for k in range(1, n + 1):
        key = f'"{cat}|{k}"'
        found = f"MATCH({key},'Scope Terms'!$E$5:$E${ST_LAST},0)"
        pr.cell(row=r, column=1, value=f'=IF(ISNA({found}),"",{k}&".")').alignment = Alignment(horizontal="right", vertical="top")
        merge_val(pr, f"B{r}:F{r}", f"=IFERROR(INDEX('Scope Terms'!$C$5:$C${ST_LAST},{found}),\"\")", align=wrap)
        window = est_lens[k - 1:k + 1] or [60]
        lines = max(1, math.ceil(max(window) / TEXTW))
        pr.row_dimensions[r].height = 12.75 * lines + 1.5
        r += 1
    r += 1

pad = lambda lst, n: lst + [60] * (n - len(lst))
block("INCLUSIONS", "Inclusion", len(INC) + 2, pad(TEXT_LEN["Inclusion"], len(INC) + 2))
block("EXCLUSIONS", "Exclusion", len(EXC) + 2, pad(TEXT_LEN["Exclusion"], len(EXC) + 2))
block("QUALIFICATIONS", "Qualification", len(QUAL) + 2, pad(TEXT_LEN["Qualification"], len(QUAL) + 2))
block("TERMS & CONDITIONS", "Term", len(TERMS) + 2, pad(TEXT_LEN["Term"], len(TERMS) + 2))

merge_val(pr, f"A{r}:F{r}", "We appreciate the opportunity and look forward to working with you on this project. Please call with any questions.", align=wrap)
pr.row_dimensions[r].height = 28
r += 2
hdr(pr, r, ["", "ACCEPTANCE", "", "", "", ""])
pr.merge_cells(f"A{r}:F{r}"); pr[f"A{r}"].value = "ACCEPTANCE"
r += 1
merge_val(pr, f"A{r}:F{r}", "The above prices, scope, and conditions are accepted. Four Corners Building Supply is authorized to proceed as specified.", align=wrap)
pr.row_dimensions[r].height = 28
r += 2
for lab_l, lab_r in (("Customer:", "Four Corners Building Supply"), ("Signature:", "Signature:"), ("Printed Name:", "Printed Name:"), ("Title:", "Title:"), ("Date:", "Date:")):
    pr.cell(row=r, column=1, value=lab_l if lab_l != "Customer:" else "").font = f_bold
    if lab_l == "Customer:":
        merge_val(pr, f"A{r}:C{r}", f'=IF({J["gc"]}="","Customer",{J["gc"]})', f_bold)
        merge_val(pr, f"D{r}:F{r}", lab_r, f_bold)
    else:
        pr.merge_cells(f"A{r}:B{r}"); pr[f"A{r}"].value = lab_l; pr[f"A{r}"].font = f_bold
        pr[f"C{r}"].border = Border(bottom=thin)
        pr[f"D{r}"].value = lab_r; pr[f"D{r}"].font = f_bold
        pr.merge_cells(f"E{r}:F{r}"); pr[f"E{r}"].border = Border(bottom=thin); pr[f"F{r}"].border = Border(bottom=thin)
    pr.row_dimensions[r].height = 22
    r += 1
PR_LAST = r
pr.print_area = f"A1:F{PR_LAST}"
pr.page_setup.orientation = "portrait"
pr.page_setup.fitToWidth = 1
pr.page_setup.fitToHeight = 0
pr.sheet_properties.pageSetUpPr.fitToPage = True
pr.page_margins.left = pr.page_margins.right = 0.5
pr.page_margins.top = pr.page_margins.bottom = 0.6
pr.oddFooter.center.text = "&8Page &P of &N"
pr.oddFooter.right.text = "&8Four Corners Building Supply"
pr.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Area Breakout (customer-facing)
ab = wb.create_sheet("Area Breakout")
ab["A1"] = f'={J["co_name"]}&"  -  AREA BREAKOUT"'; ab["A1"].font = f_title
ab["A2"] = f'="Job: "&{J["job"]}&"   |   Proposal "&{J["prop_no"]}&"   |   "&IF({J["tax_on"]}="Yes",TEXT({J["tax"]},"0.0%")&" sales tax included in all material prices","Sales tax not included")'
ab["A2"].font = f_bold
ab["A3"] = "Base bid only. This is a lump-sum bid, broken out for the GC's convenience."
ab["A3"].font = Font(name=F, size=9, italic=True)
ab.column_dimensions["A"].width = 40
area_cols = list(range(2, 10))  # B..I
hdr(ab, 5, ["CSI SECTION"] + [""] * 8 + ["Unassigned", "TOTAL"])
for i, col in enumerate(area_cols):
    c = ab.cell(row=5, column=col, value=f'=IF({AREA_CELLS[i]}="","",{AREA_CELLS[i]})')
    ab.column_dimensions[L(col)].width = 14
ab.column_dimensions["J"].width = 13
ab.column_dimensions["K"].width = 15
ab.row_dimensions[5].height = 30
r = 6
sec_r0 = r
for k in range(1, NSUM + 1):
    m_ = f"MATCH({k},Lists!$D$2:$D${NSEC + 1},0)"
    secv = f"INDEX(Lists!$B$2:$B${NSEC + 1},{m_})"
    ab.cell(row=r, column=1, value=f'=IFERROR({secv},"")').font = f_norm
    for i, col in enumerate(area_cols):
        a = AREA_CELLS[i]
        c = ab.cell(row=r, column=col, value=f'=IF($A{r}="","",IF({a}="",0,SUMIFS({rng("R")},{rng("C")},$A{r},{rng("D")},{a},{rng("B")},"Base")))')
        c.number_format = CUR0
    c = ab.cell(row=r, column=11, value=f'=IFERROR(INDEX(Lists!$C$2:$C${NSEC + 1},{m_}),"")'); c.number_format = CUR0; c.font = f_bold
    c = ab.cell(row=r, column=10, value=f'=IF($A{r}="","",K{r}-SUM(B{r}:I{r}))'); c.number_format = CUR0
    r += 1
sec_r1 = r - 1
ab.cell(row=r, column=1, value="TOTAL BY AREA").font = f_bold
for col in range(2, 12):
    c = ab.cell(row=r, column=col, value=f"=SUM({L(col)}{sec_r0}:{L(col)}{sec_r1})")
    c.number_format = CUR0; c.font = f_bold; c.fill = fill_tot; c.border = Border(top=thin)
r += 2
hdr(ab, r, ["AREA", "MATERIAL", "LABOR", "TOTAL"])
r += 1
a_r0 = r
for i in range(8):
    a = AREA_CELLS[i]
    ab.cell(row=r, column=1, value=f'=IF({a}="","",{a})')
    for col, src in ((2, "P"), (3, "Q")):
        c = ab.cell(row=r, column=col, value=f'=IF({a}="",0,SUMIFS({rng(src)},{rng("D")},{a},{rng("B")},"Base"))'); c.number_format = CUR0
    c = ab.cell(row=r, column=4, value=f"=B{r}+C{r}"); c.number_format = CUR0
    r += 1
ab.cell(row=r, column=1, value="Unassigned / Misc.")
c = ab.cell(row=r, column=2, value=f'=SUMIFS({rng("P")},{rng("B")},"Base")-SUM(B{a_r0}:B{r - 1})'); c.number_format = CUR0
c = ab.cell(row=r, column=3, value=f'=SUMIFS({rng("Q")},{rng("B")},"Base")-SUM(C{a_r0}:C{r - 1})'); c.number_format = CUR0
c = ab.cell(row=r, column=4, value=f"=B{r}+C{r}"); c.number_format = CUR0
r += 1
ab.cell(row=r, column=1, value="JOB TOTALS (before bond)").font = f_bold
for col in (2, 3, 4):
    c = ab.cell(row=r, column=col, value=f"=SUM({L(col)}{a_r0}:{L(col)}{r - 1})")
    c.number_format = CUR0; c.font = f_bold; c.fill = fill_tot; c.border = Border(top=thin)
ab.page_setup.orientation = "landscape"
ab.page_setup.fitToWidth = 1; ab.page_setup.fitToHeight = 0
ab.sheet_properties.pageSetUpPr.fitToPage = True
ab.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Bid Summary (customer-facing detail)
bs = wb.create_sheet("Bid Summary")
bs["A1"] = f'={J["co_name"]}&"  -  BID SUMMARY"'; bs["A1"].font = f_title
bs["A2"] = f'="Job: "&{J["job"]}&"   |   Proposal "&{J["prop_no"]}&"   |   Plans dated "&{dt("plans")}'
bs["A2"].font = f_bold
bs["A3"] = f'="Part of the proposal. Unit prices include material"&IF({J["tax_on"]}="Yes"," (with "&TEXT({J["tax"]},"0.0%")&" sales tax)","")&" and installation labor where applicable. Lines marked Alt or Optional are NOT in the base bid."'
bs["A3"].font = Font(name=F, size=9, italic=True)
bcols = [("No.", 6, "A"), ("Group", 9, "B"), ("CSI Section", 27, "C"), ("Area", 15, "D"), ("Mark", 12, "E"),
         ("Description", 52, "F"), ("Qty", 11, "G"), ("UOM", 6, "H"), ("Unit Price", 12, "O"),
         ("Total Price", 14, "R"), ("Notes", 20, "S")]
hdr(bs, 5, [b[0] for b in bcols])
for i, (_, w, _) in enumerate(bcols, 1):
    bs.column_dimensions[L(i)].width = w
for rr in range(6, 6 + (LAST - FIRST + 1)):
    k = rr - 5
    hc = bs.cell(row=rr, column=13, value=f"=IFERROR(MATCH({k},{rng('W')},0),0)")
    hc.font = Font(name=F, size=8, color="808080")
    m_ = f"$M{rr}"
    for i, (_, _, src) in enumerate(bcols, 1):
        if src == "A":
            v = f'=IF({m_}=0,"",{k})'
        else:
            v = f'=IF({m_}=0,"",IF(INDEX({rng(src)},{m_})="","",INDEX({rng(src)},{m_})))'
        c = bs.cell(row=rr, column=i, value=v); c.font = f_norm
        if src in ("O", "R"): c.number_format = CUR
        if src == "G": c.number_format = '#,##0.00'
        if src in ("F", "S"): c.alignment = wrap
bs.freeze_panes = "A6"
bs.column_dimensions["M"].hidden = True
bs.page_setup.orientation = "landscape"
bs.page_setup.fitToWidth = 1; bs.page_setup.fitToHeight = 0
bs.sheet_properties.pageSetUpPr.fitToPage = True
bs.print_title_rows = "5:5"
bs.sheet_view.showGridLines = False

# ---------------------------------------------------------------- Bid Comparison (VE / revision)
bc = wb.create_sheet("Bid Comparison")
bc["A1"] = "BID COMPARISON - VE / REVISION DELTAS"; bc["A1"].font = f_title
bc["A2"] = "Paste the prior bid's section totals into the yellow columns (or a competitor's number) to show the GC the deltas, like Carter's Pepper Hall VE sheet."
bc["A2"].font = Font(name=F, size=9, italic=True)
bc["A3"] = "Prior bid label:"; bc["A3"].font = f_bold
inp(bc["B3"], "Rev 0 - [date]")
bc.column_dimensions["A"].width = 40
for col in range(2, 12):
    bc.column_dimensions[L(col)].width = 14
hdr(bc, 5, ["CSI SECTION", "CURRENT MAT", "CURRENT LABOR", "CURRENT TOTAL", "PRIOR MAT", "PRIOR LABOR", "PRIOR TOTAL",
            "DELTA MAT", "DELTA LABOR", "DELTA TOTAL", "% CHANGE"])
bc.row_dimensions[5].height = 30
r = 6
for k in range(1, NSEC + 1):
    lr = k + 1
    show = f'Lists!$B${lr}<>""'
    bc.cell(row=r, column=1, value=f'=IF({show},Lists!$B${lr},"")')
    for col, src in ((2, "P"), (3, "Q")):
        c = bc.cell(row=r, column=col, value=f'=IF({show},SUMIFS({rng(src)},{rng("C")},Lists!$B${lr},{rng("B")},"Base"),0)'); c.number_format = CUR0
    bc.cell(row=r, column=4, value=f"=B{r}+C{r}").number_format = CUR0
    inp(bc.cell(row=r, column=5), None, CUR0); inp(bc.cell(row=r, column=6), None, CUR0)
    bc.cell(row=r, column=7, value=f"=N(E{r})+N(F{r})").number_format = CUR0
    bc.cell(row=r, column=8, value=f"=B{r}-N(E{r})").number_format = CUR0
    bc.cell(row=r, column=9, value=f"=C{r}-N(F{r})").number_format = CUR0
    bc.cell(row=r, column=10, value=f"=D{r}-G{r}").number_format = CUR0
    bc.cell(row=r, column=11, value=f'=IF(G{r}=0,"",J{r}/G{r})').number_format = PCT
    r += 1
br_last = r - 1
bc.cell(row=r, column=1, value="Bond").font = f_norm
bc.cell(row=r, column=4, value=f"={BOND_REF}").number_format = CUR0
inp(bc.cell(row=r, column=7), None, CUR0)
bc.cell(row=r, column=10, value=f"=D{r}-N(G{r})").number_format = CUR0
r += 1
bc.cell(row=r, column=1, value="TOTALS").font = f_bold
for col in range(2, 11):
    c = bc.cell(row=r, column=col, value=f"=SUM({L(col)}6:{L(col)}{r - 1})")
    c.number_format = CUR0; c.font = f_bold; c.fill = fill_tot; c.border = Border(top=thin)
c = bc.cell(row=r, column=11, value=f'=IF(G{r}=0,"",J{r}/G{r})'); c.number_format = PCT; c.font = f_bold; c.fill = fill_tot
bc.freeze_panes = "B6"

# ---------------------------------------------------------------- Instructions content
ws0.column_dimensions["A"].width = 4
ws0.column_dimensions["B"].width = 26
ws0.column_dimensions["C"].width = 100
ws0["B1"] = "FOUR CORNERS COMMERCIAL QUOTE TEMPLATE"; ws0["B1"].font = f_title
ws0["B2"] = "For multifamily and commercial Div 6 / 8 / 10 bids (trim, doors, frames, hardware, specialties). Built from the Carter Lumber proposal format GCs already know."
ws0["B2"].font = Font(name=F, size=10, italic=True)
steps = [
    ("HOW TO USE", None),
    ("1. Save As", "Save a copy per job: [Job Name] - [GC] - Proposal [#].xlsx. Never type in the master template."),
    ("2. Job Info", "Fill every yellow cell: GC, job, plans date, estimator, tax rate (VERIFY the county), margins, bond, dates, and area names."),
    ("3. Line Detail", "One row per takeoff item. Pick a Bid Group (Base / Alt 1-5 / Optional / Excluded), CSI Section, Area, Qty, UOM, and enter COST. Sell price, tax, and extensions calculate. Red cells = missing required info."),
    ("", "Per-line margin overrides go in the two 'Override' columns (blank = Job Info default). Delete the 4 blue-italic EXAMPLE rows first."),
    ("4. Alternates", "Describe each alternate. Amounts total automatically from Line Detail lines tagged Alt 1-5. Use negative quantities for deducts."),
    ("5. Scope Terms", "Set each inclusion, exclusion, qualification, and term to Yes/No for this job. Only 'Yes' items print, auto-numbered. Edit wording as needed."),
    ("6. Review", "Check the JOB MARGIN CHECK block at the bottom of Job Info: unassigned $ and missing Bid Group count must both be zero. Look at GP% and labor %."),
    ("7. Send", "Print to PDF: Proposal + Area Breakout + Bid Summary (in that order). NEVER send Line Detail or Job Info; they show cost and margin."),
    ("", "Bid Summary's print area grows with the number of lines automatically in Excel. If it ever prints blank pages, select the filled rows and use Page Layout > Print Area > Set Print Area."),
    ("8. VE / Revisions", "Before re-pricing, copy the current section totals into 'Bid Comparison' prior columns, then make changes. The sheet shows the deltas for the GC."),
    ("", None),
    ("COLOR KEY", None),
    ("Yellow + blue text", "Input - type here."),
    ("Green text", "Pulls from another sheet - don't overwrite unless you mean to."),
    ("Black text", "Formula - don't touch."),
    ("Blue italic", "Example data - delete before use."),
    ("", None),
    ("SHEETS", None),
    ("Proposal", "CUSTOMER. Cover letter, CSI summary, bond, base total, alternates, inclusions/exclusions/qualifications/terms, signature block."),
    ("Area Breakout", "CUSTOMER. Section x area matrix plus material/labor by area (Carter's 'Apartments / Clubhouse / Pool Cabana' pages)."),
    ("Bid Summary", "CUSTOMER. Numbered line items with unit and extended price, compacted automatically (no blank rows, Excluded lines removed)."),
    ("Bid Comparison", "CUSTOMER (optional). VE or revision delta sheet."),
    ("Line Detail", "INTERNAL. Takeoff with cost, margin, and sell."),
    ("Job Info", "INTERNAL. Inputs and margin check."),
    ("Scope Terms", "INTERNAL. Library of scope language."),
    ("Alternates / Lists", "INTERNAL. Supporting tables. Add custom CSI sections in the yellow rows on Lists."),
    ("", None),
    ("BEFORE FIRST USE", "Have Wesley (company attorney) review the Terms & Conditions and Qualifications on 'Scope Terms', especially the finance charge, lien, warranty, and escalation language."),
]
r = 4
for a, b in steps:
    ca = ws0.cell(row=r, column=2, value=a or None)
    if b is None and a:
        ca.font = f_sec
    else:
        ca.font = f_bold
        cb = ws0.cell(row=r, column=3, value=b); cb.font = f_norm; cb.alignment = wrap
        if b and len(b) > 100:
            ws0.row_dimensions[r].height = 28
    r += 1
ws0["B17"].fill = fill_in; ws0["B17"].font = f_input
ws0["B18"].font = f_link
ws0["B20"].font = f_ex
ws0.sheet_view.showGridLines = False

# order sheets: Instructions, Job Info, Line Detail, Alternates, Scope Terms, Proposal, Area Breakout, Bid Summary, Bid Comparison, Lists
order = ["Instructions", "Job Info", "Line Detail", "Alternates", "Scope Terms", "Proposal", "Area Breakout", "Bid Summary", "Bid Comparison", "Lists"]
wb._sheets = [wb[n] for n in order]
for n in ("Proposal", "Area Breakout", "Bid Summary", "Bid Comparison"):
    wb[n].sheet_properties.tabColor = "2E75B6"
for n in ("Job Info", "Line Detail", "Alternates", "Scope Terms", "Lists"):
    wb[n].sheet_properties.tabColor = "BF9000"

# ensure Arial everywhere
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != F:
                ft = c.font
                c.font = Font(name=F, size=ft.size, bold=ft.bold, italic=ft.italic, color=ft.color)
from openpyxl.workbook.defined_name import DefinedName
for ws_ in (ji, ld, st, al, bc, ws0):
    ws_.page_setup.orientation = "landscape"
    ws_.page_setup.fitToWidth = 1; ws_.page_setup.fitToHeight = 0
    ws_.sheet_properties.pageSetUpPr.fitToPage = True
ld.print_title_rows = "5:5"
ld.defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area", attr_text=f"OFFSET({LD}!$A$1,0,0,{FIRST - 1}+MAX({LD}!$W${FIRST}:$W${LAST}),24)")
bs.defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area", attr_text=f"OFFSET('Bid Summary'!$A$1,0,0,5+MAX({LD}!$W${FIRST}:$W${LAST}),11)")
wb.active = 0
wb.save(OUT)
print("saved", OUT)
